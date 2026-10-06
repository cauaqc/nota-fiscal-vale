"""
Funções comuns aos fluxos de locação e serviço: planilha de controle
(Billing), aviso à API e o laço que processa uma lista de notas. A extração do
PDF da locação está em locacao/lancar_locacao.py.
"""

import os
import re
import unicodedata
from datetime import date

import openpyxl
import pdfplumber

from api.api_client import (
    STATUS_DUPLICADO,
    STATUS_ERRO,
    STATUS_OK,
    NotasValeApiError,
    marcar_nota_processada,
)

DIAS_VENCIMENTO = 30  # regra: vencimento = emissão + 30 dias corridos


def extrair_frs_rf_servico(caminho_pdf):
    """FRS e RF da nota de serviço, escritos na descrição: 'FRS:1007520374 FR:6203204908' (o RF vem como 'FR' ou 'RF')."""
    with pdfplumber.open(caminho_pdf) as pdf:
        texto = " ".join((p.extract_text() or "") for p in pdf.pages)
    frs = re.search(r"FRS\s*:?\s*(\d+)", texto)
    rf = re.search(r"\b(?:FR|RF)\s*:?\s*(\d+)", texto)
    return (frs.group(1) if frs else None, rf.group(1) if rf else None)


MESES = [
    "janeiro", "fevereiro", "marco", "abril", "maio", "junho",
    "julho", "agosto", "setembro", "outubro", "novembro", "dezembro",
]


def _sem_acento(texto):
    return unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode().lower()


def localizar_planilha_do_mes(pasta, hoje=None):
    """
    Acha em `pasta` a planilha Billing do mês de CONSUMO, que é o mês anterior
    ao atual (as notas emitidas em outubro são do consumo de setembro). O nome
    do arquivo precisa ter o nome do mês e o ano, ex.: 'Billing Vale Setembro- 2026.xlsx'.
    Levanta FileNotFoundError se não achar exatamente um arquivo.
    """
    hoje = hoje or date.today()
    mes = hoje.month - 1 or 12
    ano = hoje.year if hoje.month > 1 else hoje.year - 1
    nome_mes = MESES[mes - 1]

    candidatos = [
        nome
        for nome in os.listdir(pasta)
        if nome.lower().endswith(".xlsx")
        and not nome.startswith("~$")  # arquivo de trava do Excel aberto
        and nome_mes in _sem_acento(nome)
        and str(ano) in nome
    ]
    if len(candidatos) != 1:
        raise FileNotFoundError(
            f"esperava 1 planilha de {nome_mes}/{ano} em {pasta}, achei {len(candidatos)}: {candidatos}"
        )
    return os.path.join(pasta, candidatos[0])


def _coluna(cabecalhos, predicado):
    achadas = [i for i, c in enumerate(cabecalhos, start=1) if c and predicado(str(c).upper())]
    return achadas[0] if len(achadas) == 1 else None


def _colunas_da_aba(ws, linhas_cabecalho=5):
    """
    Procura o cabeçalho nas primeiras linhas da aba (na Billing real ele fica
    na linha 2; a 1 tem os SUBTOTAL). Retorna (linha_cabecalho, frs, rf,
    chamado, nf) ou None se a aba não tiver as 4 colunas.
    """
    for linha in range(1, min(linhas_cabecalho, ws.max_row) + 1):
        cabecalhos = [c.value for c in ws[linha]]
        colunas = (
            _coluna(cabecalhos, lambda t: "(FRS)" in t),
            _coluna(cabecalhos, lambda t: "(RF)" in t),
            _coluna(cabecalhos, lambda t: "CHAMADO" in t),
            _coluna(cabecalhos, lambda t: t.strip() == "NF"),
        )
        if all(colunas):
            return (linha, *colunas)
    return None


def _como_numero(valor):
    texto = str(valor).strip()
    return int(texto) if texto.isdigit() else texto


def registrar_na_planilha(caminho_planilha, frs, rf, numero_nf, processo):
    """
    Acha, em todas as abas que têm as colunas 'Nº FOLHA DE REGISTRO (FRS)',
    'Nº RELATÓRIO DE FATURAMENTO (RF)', 'N° CHAMADO' e 'NF', a linha cujo FRS e
    RF batem com os da nota e preenche 'N° CHAMADO' (nº do processo na Vale) e
    'NF'. Retorna 'aba!linha'.

    Falha (LookupError / ValueError) sem escrever nada se: o FRS/RF não existir
    ou existir em mais de uma linha, ou se o 'N° CHAMADO' da linha já estiver
    preenchido com outro valor (nunca sobrescreve).
    """
    if not frs or not rf:
        raise LookupError("FRS/RF da nota não encontrados")

    wb = openpyxl.load_workbook(caminho_planilha)
    achadas = []
    for ws in wb.worksheets:
        colunas = _colunas_da_aba(ws)
        if not colunas:
            continue
        linha_cabecalho, col_frs, col_rf, col_chamado, col_nf = colunas
        achadas += [
            (ws, r, col_chamado, col_nf)
            for r in range(linha_cabecalho + 1, ws.max_row + 1)
            if str(ws.cell(r, col_frs).value).strip() == str(frs)
            and str(ws.cell(r, col_rf).value).strip() == str(rf)
        ]
    if len(achadas) != 1:
        raise LookupError(f"FRS {frs} + RF {rf} casam com {len(achadas)} linha(s) da planilha (esperado 1)")
    ws, linha, col_chamado, col_nf = achadas[0]
    onde = f"{ws.title}!{linha}"

    atual = ws.cell(linha, col_chamado).value
    if atual not in (None, "") and str(atual).strip() != str(processo):
        raise ValueError(f"{onde} já tem N° CHAMADO = {atual} (não sobrescrevo)")

    ws.cell(linha, col_chamado).value = _como_numero(processo)
    ws.cell(linha, col_nf).value = _como_numero(numero_nf)
    # Formato "0": número inteiro sem separador de milhar (alguns
    # visualizadores mostram o "Geral" como 9,053,422).
    ws.cell(linha, col_chamado).number_format = "0"
    ws.cell(linha, col_nf).number_format = "0"
    wb.save(caminho_planilha)
    return onde


class EnvioIncerto(Exception):
    """Falhou depois do clique em 'Ingressar': a nota pode ter sido criada."""


class DadosInvalidos(Exception):
    """O PDF não trouxe todos os campos necessários: a nota é pulada."""


def anotar_na_planilha(r, caminho_planilha):
    """Preenche N° CHAMADO e NF na planilha. Falha aqui não desfaz o envio, só avisa."""
    try:
        onde = registrar_na_planilha(
            caminho_planilha, r.get("frs"), r.get("rf"), r["numero_nf"], r["processo_vale"]
        )
        print(f"   📊 planilha: {onde} ← N° CHAMADO {r['processo_vale']}, NF {r['numero_nf']}")
    except Exception as e:  # noqa: BLE001
        print(f"   ⚠️  NÃO preenchi a planilha (NF {r['numero_nf']}, processo #{r['processo_vale']}): {e}")


def avisar_api(r, status):
    """POST do resultado. Falha aqui não derruba o loop, mas é avisada."""
    try:
        marcar_nota_processada(r["numero_nf"], r["coligada"], status)
        print(f"   📮 API avisada: NF {r['numero_nf']} = {status}")
        return True
    except NotasValeApiError as e:
        print(f"   ⚠️  NÃO consegui avisar a API (NF {r['numero_nf']} = {status}): {e}")
        return False


def processar_notas(context, notas, kind, lancar_uma, enviar, planilha):
    """
    Laço comum aos dois fluxos. Para cada nota: lança no portal (`lancar_uma`),
    anota o nº do processo na planilha, avisa a API e fecha a aba.
    Uma nota que falha não derruba as demais. Retorna [(numero_nf, kind, texto)].
    """
    resumo = []
    for r in notas:
        numero = r["numero_nf"]
        print(f"=== NF {numero} ({kind}) ===")
        status_api, texto = None, ""
        try:
            status_api, texto = lancar_uma(context, r, enviar)
        except EnvioIncerto as e:
            print(f"   ❗ falha DEPOIS do clique em 'Ingressar': {e}")
            resumo.append((numero, kind, "VERIFICAR NO PORTAL (envio incerto, API não avisada)"))
            continue
        except Exception as e:  # noqa: BLE001
            status_api, texto = STATUS_ERRO, f"erro: {str(e).splitlines()[0]}"
            print(f"   ❌ {texto}")

        if r.get("processo_vale"):
            print(f"   🔢 processo da Vale: #{r['processo_vale']}")
            anotar_na_planilha(r, planilha)

        if enviar and status_api:
            if not avisar_api(r, status_api):
                texto += " | API NÃO avisada"

        # Aba da nota: fecha só se foi enviada e o nº do processo foi lido
        # (evita acumular uma aba por nota). Em teste, erro ou envio incerto
        # a aba fica aberta para você conferir.
        aba = r.pop("_aba", None)
        if aba and enviar and status_api in (STATUS_OK, STATUS_DUPLICADO) and r.get("processo_vale"):
            aba.close()

        resumo.append((numero, kind, texto))
    return resumo
